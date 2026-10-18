"""
OutcomeX — Excel Template Generation & Adapter Test Suite
Tests openpyxl adapter fidelity, cell positioning, formula preservation/repair,
and verifies that the reference template remains pristine.
"""

import os
import unittest
import openpyxl
from scripts.excel_generator import (
    Adapter60_40,
    Adapter30_20,
    generate_workbook,
    compute_sha256,
    REFERENCE_WORKBOOK_PATH,
    EXPORTS_DIR,
)

ORIGINAL_TEMPLATE_SHA256 = "497f4d82b76337c5826027f606e635489777ee9f8b95c26e878e84c45a6ad477"

class TestExcelGeneration(unittest.TestCase):
    def setUp(self):
        os.makedirs(EXPORTS_DIR, exist_ok=True)
        self.sample_course = {
            "course_code": "25UCS623T_TEST",
            "course_name": "Test Course in Data Science",
            "department": "CSE",
            "year_semester": "3rd yr/5th sem",
            "academic_year": "2025-26",
            "po_results": [
                {"po_code": "PO1", "achieved_attainment": 2.67},
                {"po_code": "PO2", "achieved_attainment": 2.67},
            ]
        }
        self.sample_students = [
            {
                "roll_no": "101",
                "student_name": "TEST STUDENT ONE",
                "marks": {
                    "mt1_co1": 9.5,
                    "mt1_co2": 8.0,
                    "mt1_co3": 7.5,
                    "mt2_co4": 9.0,
                    "mt2_co5": 8.5,
                    "mt2_co6": 10.0,
                    "ese_co1": 7.0,
                    "ese_co2": 8.0,
                    "ese_co3": 9.0,
                    "ese_co4": 6.5,
                    "ese_co5": 8.0,
                    "ese_co6": 7.5,
                    "ta": 23.0,
                }
            },
            {
                "roll_no": "102",
                "student_name": "TEST STUDENT TWO",
                "marks": {
                    "mt1_co1": "A", # Absent
                    "mt1_co2": 5.0,
                    "ta": 20.0,
                }
            }
        ]

    def test_01_template_integrity_unmodified(self):
        """Verifies reference template file exists and has not been corrupted or modified."""
        self.assertTrue(os.path.exists(REFERENCE_WORKBOOK_PATH))
        current_sha = compute_sha256(REFERENCE_WORKBOOK_PATH)
        self.assertEqual(current_sha, ORIGINAL_TEMPLATE_SHA256, "Reference template SHA-256 must match exactly")

    def test_02_adapter_60_40_generation(self):
        """Tests Adapter60_40 correctly populates metadata, marks, and preserves sheets."""
        output_path = os.path.join(EXPORTS_DIR, "test_output_60_40.xlsx")
        adapter = Adapter60_40()
        res = adapter.generate_report(self.sample_course, self.sample_students, output_path=output_path)

        self.assertTrue(res["success"])
        self.assertTrue(os.path.exists(output_path))
        self.assertGreater(os.path.getsize(output_path), 50000)

        # Inspect generated workbook
        wb = openpyxl.load_workbook(output_path, data_only=False)
        self.assertIn("For 60_40 (final)", wb.sheetnames)
        ws = wb["For 60_40 (final)"]

        # Check metadata
        self.assertEqual(ws["O3"].value, "CSE")
        self.assertEqual(ws["O4"].value, "Test Course in Data Science")
        self.assertEqual(ws["O5"].value, "25UCS623T_TEST")

        # Check student 1 (row 29)
        self.assertEqual(ws["A29"].value, 101.0)
        self.assertEqual(ws["B29"].value, "TEST STUDENT ONE")
        self.assertEqual(ws["C29"].value, 9.5)
        self.assertEqual(ws["D29"].value, 8.0)
        self.assertEqual(ws["E29"].value, 7.5)
        self.assertEqual(ws["T29"].value, 23.0)

        # Check student 2 (row 30 with absent string)
        self.assertEqual(ws["A30"].value, 102.0)
        self.assertEqual(ws["B30"].value, "TEST STUDENT TWO")
        self.assertEqual(ws["C30"].value, "A")
        self.assertEqual(ws["D30"].value, 5.0)

        # Check page setup orientation
        self.assertEqual(ws.page_setup.orientation, ws.ORIENTATION_LANDSCAPE)
        wb.close()

    def test_03_adapter_30_20_generation_and_repairs(self):
        """Tests Adapter30_20 fixes #REF! in Y7:AB7, E97 typo in Y4:AB4, and G27/H27 labels."""
        output_path = os.path.join(EXPORTS_DIR, "test_output_30_20.xlsx")
        adapter = Adapter30_20()
        res = adapter.generate_report(self.sample_course, self.sample_students, output_path=output_path)

        self.assertTrue(res["success"])
        self.assertTrue(os.path.exists(output_path))

        wb = openpyxl.load_workbook(output_path, data_only=False)
        self.assertIn("For 30_20 (Final)", wb.sheetnames)
        ws = wb["For 30_20 (Final)"]

        # 1. Verify headers G27 and H27 are repaired to CO3 and CO4
        self.assertEqual(ws["G27"].value, "CO3")
        self.assertEqual(ws["H27"].value, "CO4")

        # 2. Verify Y4:AB4 references level row 98 ($C$98 etc.)
        self.assertIn("$C$98", ws["Y4"].value)
        self.assertNotIn("E97", ws["Y4"].value)

        # 3. Verify Y7:AB7 formulas are reconstructed without #REF!
        self.assertNotIn("#REF!", ws["Y7"].value)
        self.assertIn("$K$94", ws["Y7"].value)
        self.assertIn("$K$98", ws["Y7"].value)

        # 4. Check student row 29
        self.assertEqual(ws["A29"].value, 101.0)
        self.assertEqual(ws["B29"].value, "TEST STUDENT ONE")
        wb.close()

    def test_04_invalid_scheme_handling(self):
        """Verifies invalid scheme raises ValueError."""
        with self.assertRaises(ValueError):
            generate_workbook("invalid_scheme_99")

    def test_05_original_template_unaltered_after_generations(self):
        """Verifies template checksum remains completely identical after multiple generations."""
        current_sha = compute_sha256(REFERENCE_WORKBOOK_PATH)
        self.assertEqual(current_sha, ORIGINAL_TEMPLATE_SHA256)


if __name__ == "__main__":
    unittest.main()
